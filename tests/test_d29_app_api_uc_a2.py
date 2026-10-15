# -*- coding: utf-8 -*-
"""ÖH Dalga-29 / kulvar ``app-api-uc``, aşama 2 — uç bekçileri.

Bu dosyadaki her sınıf, ``docs/scofield/raw/OH/D29/app-api-uc/a2/once.json``
içinde TABAN AĞACINDA (``4dc63c11``) ölçülmüş bir kusuru kilitler. Her
sınıfın künyesinde ölçümün kendisi durur; test o ölçümün geri gelmesini
engeller. Ters-mutasyon kanıtı: ``a2/mutasyon.json``.

Kalemler:

* **P16B4-20** — ``/api/export-pdf/<tip>``: bilinmeyen tür sessizce
  ``complete``e düşüyordu (``HAYALET`` -> 200 + 7.109 B +
  ``motor_complete_*.pdf``).
* **P16B4-23** — ``/api/export-xlsx``: ``NaN``/``Inf`` hücreler sessizce
  BOŞ iniyordu (Excel'de "veri yok" demek).
* **P16B4-24** — ``/api/export-xlsx``: ``rows`` metin olduğunda karakter
  başına satır üretiliyordu; ``rows: 42`` ham HTTP 500 veriyordu.
* **P16B3-33** — ``/api/six-dof-analysis``: ``fin_count = 3.5`` HTTP 200
  ve ``fin_count = 3`` ile aynı uzunlukta yanıt (çözücü sessizce kırpıyor).
* **P16B3-32** — ters birim şüphesi: beş gövdeden dördü ZATEN kapılıydı,
  açık kalan tek ayak ``/api/flow-analysis`` idi (30 m boğaz -> HTTP 200).
* **P7B1-20** — ``/calculate``: ``include_plots='false'`` metni ``bool()``
  ile ``True``ya dönüyordu (2.678.462 B + tam figür seti).
* **P7B1-21** — ``hrma/app.py:_mm_to_m`` ölü yardımcıydı (0 çağrı) ve
  künyesi yürürlükten kalkmış sözleşmeyi anlatıyordu.
* **S1-05** — ihracat ailesinde HTTP sınıfı: BOŞ gövdede 9 uç 422 verirken
  ``/api/export-stl`` 400 veriyordu.
"""
import ast
import contextlib
import io
import json
import math
import os

import pytest

from hrma.app import app
from tests.support.kanit import kanit_gerekli


@pytest.fixture
def client():
    app.config['TESTING'] = True
    with app.test_client() as c:
        yield c


def _post(client, yol, govde=None, ham=None):
    with contextlib.redirect_stdout(io.StringIO()):
        with contextlib.redirect_stderr(io.StringIO()):
            if ham is not None:
                return client.post(yol, data=ham,
                                   content_type='application/json')
            return client.post(yol, json=govde)


def _xlsx_hucreler(bayt):
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(bayt))
    ws = wb[wb.sheetnames[0]]
    return [[c.value for c in satir] for satir in ws.iter_rows()]


# ===========================================================================
# P16B4-20 — RAPOR TÜRÜNÜN ADI UYDURULMAZ
# ===========================================================================
PDF_GOVDE = {'motor_data': {'motor_name': 'FUZZBASE', 'thrust': 1000.0,
                            'chamber_pressure': 20.0, 'burn_time': 10.0},
             'analysis_results': {}}


class TestP16B420RaporTuru:
    """ÖLÇÜLDÜ (``once.json`` K1): ``/api/export-pdf/HAYALET`` ve
    ``/api/export-pdf/analysis`` HTTP 200 + 7.109 B + ``complete``in dosya
    adıyla PDF veriyordu — ``complete`` ile aynı uzunlukta.
    """

    @pytest.mark.parametrize('tip', ['HAYALET', 'analysis', 'COMPLETE',
                                     'summary ', 'quick'])
    def test_taninmayan_tur_belge_uretmez(self, client, tip):
        r = _post(client, '/api/export-pdf/' + tip, dict(PDF_GOVDE))
        assert r.status_code == 400, (
            '%r: HTTP %d, %d B' % (tip, r.status_code, len(r.get_data())))
        j = r.get_json() or {}
        assert j.get('error') == 'unknown_report_type', j
        alan = (j.get('invalid_fields') or [{}])[0]
        assert alan.get('field') == 'report_type', j
        assert alan.get('submitted') == tip, j
        # Kabul edilen küme makine-okur biçimde YAYIMLANIR (P16B1-47).
        assert sorted(alan.get('accepted_values') or []) == [
            'complete', 'summary', 'technical'], j
        assert not r.get_data().startswith(b'%PDF-'), 'PDF üretilmiş'

    @pytest.mark.parametrize('tip', ['summary', 'technical', 'complete'])
    def test_kurulu_uc_tur_aynen_gecer(self, client, tip):
        r = _post(client, '/api/export-pdf/' + tip, dict(PDF_GOVDE))
        assert r.status_code == 200, r.get_data(as_text=True)[:200]
        assert r.get_data()[:5] == b'%PDF-'
        assert ('motor_%s_FUZZBASE.pdf' % tip) in (
            r.headers.get('Content-Disposition') or '')

    def test_kabul_kumesi_tek_kaynaktan(self):
        """Kapı ile dallanma AYNI demeti okur (Doktrin 4)."""
        from hrma.app import PDF_REPORT_TYPES
        assert PDF_REPORT_TYPES == ('summary', 'technical', 'complete')

    def test_motor_verisi_kapisi_tur_kapisindan_ONCE(self, client):
        """Sıra korunur: boş gövde + tanınmayan tür -> ``no_motor_data``.

        Mevcut bekçi ``tests/test_d18_ihracat_a1.py`` boş gövdeyi
        ``analysis`` türüyle de gönderiyor ve ``no_motor_data`` bekliyor.
        """
        r = _post(client, '/api/export-pdf/analysis', {})
        assert r.status_code == 422, r.status_code
        assert (r.get_json() or {}).get('error') == 'no_motor_data'


# ===========================================================================
# P16B4-24 — SATIR/BAŞLIK TAŞIYICISININ TİPİ
# ===========================================================================
class TestP16B424SatirTipi:
    """ÖLÇÜLDÜ (``once.json`` K2): ``rows: "abc"`` -> HTTP 200 ve sayfa
    ``[['v'], ['a'], ['b'], ['c']]``; ``rows: {"a": 1}`` -> ``[['v'],
    ['a']]`` (yalnız anahtar); ``rows: 42`` -> HTTP 500 ham çözücü hatası.
    """

    @pytest.mark.parametrize('rows,tip_adi', [
        ('abc', 'str'), ({'a': 1}, 'dict'), (42, 'int'), (3.5, 'float'),
    ])
    def test_liste_olmayan_rows_reddedilir(self, client, rows, tip_adi):
        r = _post(client, '/api/export-xlsx', {
            'filename': 'm.xlsx',
            'sheets': [{'name': 'S', 'headers': ['v'], 'rows': rows}]})
        assert r.status_code == 422, (
            '%r: HTTP %d' % (rows, r.status_code))
        j = r.get_json() or {}
        assert j.get('error') == 'invalid_rows', j
        alan = (j.get('invalid_fields') or [{}])[0]
        assert alan.get('field') == 'sheets[0].rows', j
        assert alan.get('reason') == 'not_a_list', j
        assert alan.get('received_type') == tip_adi, j

    def test_liste_olmayan_headers_reddedilir(self, client):
        r = _post(client, '/api/export-xlsx', {
            'sheets': [{'name': 'S', 'headers': 'ab', 'rows': [[1]]}]})
        assert r.status_code == 422, r.status_code
        j = r.get_json() or {}
        assert j.get('error') == 'invalid_headers', j
        assert (j.get('invalid_fields') or [{}])[0].get('field') == \
            'sheets[0].headers', j

    @pytest.mark.parametrize('rows,beklenen', [
        ([[1, 2], [3, 4]], [['a', 'b'], [1, 2], [3, 4]]),
        ([1, 2, 3], [['a'], [1], [2], [3]]),
        ([['x'], ['y']], [['a'], ['x'], ['y']]),
        (None, [['a']]),
        ([], [['a']]),
    ])
    def test_saglikli_envanter_gevsemez(self, client, rows, beklenen):
        """ÖK-22: ölçülen sağlıklı gövdeler bit-aynı kalır."""
        basliklar = ['a', 'b'] if rows and rows[0] == [1, 2] else ['a']
        beklenen = [basliklar] + beklenen[1:]
        r = _post(client, '/api/export-xlsx', {
            'filename': 'ok.xlsx',
            'sheets': [{'name': 'S', 'headers': basliklar, 'rows': rows}]})
        assert r.status_code == 200, r.get_data(as_text=True)[:200]
        assert _xlsx_hucreler(r.get_data()) == beklenen

    def test_ust_seviye_sheets_kapisi_bozulmaz(self, client):
        r = _post(client, '/api/export-xlsx', {'sheets': 'abc'})
        assert r.status_code == 422
        assert (r.get_json() or {}).get('error') == 'invalid_sheets'


# ===========================================================================
# P16B4-23 — TEMSİL EDİLEMEYEN HÜCRE
# ===========================================================================
class TestP16B423SonluOlmayanHucre:
    """ÖLÇÜLDÜ (``once.json`` K2 ``nan_inf``): ``[[NaN], [Infinity],
    [-Infinity], [1.5]]`` gövdesi HTTP 200 ve sayfa ``[['v'], [None],
    [None], [None], [1.5]]`` — üç hücre BOŞ indi, hiçbir yerde gerekçe yok.
    """

    @pytest.mark.parametrize('belirtec', ['NaN', 'Infinity', '-Infinity'])
    def test_sonlu_olmayan_hucre_dosya_uretmez(self, client, belirtec):
        ham = ('{"filename":"nan.xlsx","sheets":[{"name":"S",'
               '"headers":["v"],"rows":[[1.5],[%s]]}]}' % belirtec)
        r = _post(client, '/api/export-xlsx', ham=ham)
        assert r.status_code == 422, (
            '%s: HTTP %d, %d B' % (belirtec, r.status_code,
                                   len(r.get_data())))
        j = r.get_json() or {}
        assert j.get('error') == 'unrepresentable_cell', j
        alan = (j.get('invalid_fields') or [{}])[0]
        assert alan.get('field') == 'sheets[0].rows[1][0]', j
        assert alan.get('reason') == 'not_finite', j

    def test_sonlu_sayilar_ve_metin_aynen_gecer(self, client):
        r = _post(client, '/api/export-xlsx', {
            'filename': 'ok.xlsx',
            'sheets': [{'name': 'S', 'headers': ['v'],
                        'rows': [[1.5], [0], [-3], ['metin'], [None],
                                 [True]]}]})
        assert r.status_code == 200, r.get_data(as_text=True)[:200]
        hucreler = _xlsx_hucreler(r.get_data())
        assert hucreler[1] == [1.5] and hucreler[4] == ['metin'], hucreler


# ===========================================================================
# P16B3-33 — KANAT SAYISI TAMSAYIDIR
# ===========================================================================
SIX_DOF = {'body_diameter': 0.1, 'body_length': 2.0, 'nose_length': 0.3,
           'dry_mass': 10.0, 'propellant_mass': 3.0, 'cd0': 0.5,
           'fin_count': 4, 'fin_root_chord': 0.15, 'fin_tip_chord': 0.08,
           'fin_span': 0.1, 'fin_sweep': 0.05, 'thrust': 1500.0,
           'burn_time': 4.0, 'rail_length': 5.0,
           'launch_elevation_deg': 85.0}


class TestP16B333KanatSayisi:
    """ÖLÇÜLDÜ (``once.json`` K3): ``fin_count = 3.5`` HTTP 200 ve 44.814 B
    — ``fin_count = 3`` ile BİT-UZUNLUKTA; çözücü
    (``six_dof_trajectory.py``: ``n = int(fin_count)``) sessizce kırpıyordu.
    ``'3.5'`` metni ise ham ``int()`` hatasıyla 400 veriyordu.
    """

    @pytest.mark.parametrize('fc', [3.5, 0.4, 2.0001])
    def test_tamsayi_olmayan_kanat_sayisi_reddedilir(self, client, fc):
        r = _post(client, '/api/six-dof-analysis', dict(SIX_DOF,
                                                        fin_count=fc))
        assert r.status_code == 422, '%r: HTTP %d' % (fc, r.status_code)
        j = r.get_json() or {}
        assert j.get('error') == 'invalid_six_dof_input', j
        alanlar = {a.get('field'): a for a in (j.get('invalid_fields') or [])}
        assert 'fin_count' in alanlar, j
        assert alanlar['fin_count'].get('reason') == 'not_an_integer', j

    #: ÖH D33-2 / app-api-uc a1 (OH-D32-16) — DİZE VAKALARI BURAYA TAŞINDI.
    #: ``'3.5'`` yukarıdaki tablodan, ``'4'`` ise aşağıdaki "aynen geçer"
    #: tablosundan ÇIKARILDI. Sebep ölçülmüştür
    #: (``D33-2/app-api-uc/a1/once_gerekce_kanonu.json`` K1/K2): bu ucun
    #: kapısı (``_collect_unphysical_fields``) sayı-benzeri DİZEYİ sessizce
    #: kabul ediyordu, oysa kardeş uç ``POST /calculate`` aynı dizeye
    #: 400 ``not_a_json_number`` veriyordu — aynı tel biçimine iki hüküm.
    #: Kapı merkezin kendi yüklemine (``input_guard.is_json_number``)
    #: bağlandı. İDDİA ZAYIFLAMADI: ``'4'`` artık 200 DEĞİL, 422 alır ve
    #: reddin ADI vardır. Sayısal 4 / 3 / 0 / 6.0 aynen geçmeye devam eder.
    @pytest.mark.parametrize('fc', ['3.5', '4', '0'])
    def test_sayi_benzeri_DIZE_reddedilir(self, client, fc):
        r = _post(client, '/api/six-dof-analysis', dict(SIX_DOF,
                                                        fin_count=fc))
        assert r.status_code == 422, '%r: HTTP %d' % (fc, r.status_code)
        j = r.get_json() or {}
        alanlar = {a.get('field'): a for a in (j.get('invalid_fields') or [])}
        assert alanlar['fin_count'].get('reason') == 'not_a_json_number', j

    @pytest.mark.parametrize('fc', [4, 3, 0, 6.0])
    def test_tamsayi_kanat_sayisi_aynen_gecer(self, client, fc):
        r = _post(client, '/api/six-dof-analysis', dict(SIX_DOF,
                                                        fin_count=fc))
        assert r.status_code == 200, (
            '%r: HTTP %d %s' % (fc, r.status_code,
                                r.get_data(as_text=True)[:160]))


# ===========================================================================
# P16B3-32 — TERS BİRİM ŞÜPHESİ (açık kalan tek ayak: /api/flow-analysis)
# ===========================================================================
# ÖH D32 (OH-D31-14): ``ambient_pressure`` eklendi (zorunlu altıncı alan);
# değer ucun kanonik tablosundan, yayımlanan sayı DEĞİŞMEDİ.
FLOW = {'chamber_pressure': 20.0, 'throat_diameter': 0.03,
        'expansion_ratio': 8.0, 'gamma': 1.2, 'molecular_weight': 22.0,
        'chamber_temperature': 3000.0, 'ambient_pressure': 101325.0}


class TestP16B332TersBirim:
    """ÖLÇÜLDÜ (``once.json`` K4): beş gövdeden DÖRDÜ zaten kapılıydı
    (wall-profile Pa/bar 422, water-hammer m/mm 422, thermal_safety mm/m
    422, structural karşı örneği 422); ``/api/flow-analysis`` 30 m'lik
    boğazla HTTP 200 + 40.435 B tam akış hükmü veriyordu.
    """

    def test_flow_analysis_metre_mm_karisikligi_kapali(self, client):
        r = _post(client, '/api/flow-analysis', dict(FLOW,
                                                     throat_diameter=30.0))
        assert r.status_code == 400, (
            'HTTP %d, %d B' % (r.status_code, len(r.get_data())))
        j = r.get_json() or {}
        alan = (j.get('invalid_fields') or [{}])[0]
        assert alan.get('field') == 'throat_diameter', j
        assert alan.get('reason') == 'out_of_range', j
        assert alan.get('accepted_range')[1] == 2.0, j

    def test_saglikli_bogaz_bit_ayni(self, client):
        r = _post(client, '/api/flow-analysis', dict(FLOW))
        assert r.status_code == 200, r.get_data(as_text=True)[:200]

    def test_bant_kanonik_tablodan_gelir(self):
        """Yeni eşik İCAT EDİLMEDİ: bant deponun kendi tablosundandır."""
        from hrma.app import _ust_sinir_bandi
        from hrma.validation.motor_validation import motor_validator
        bant = _ust_sinir_bandi(('throat_diameter', 'chamber_pressure'))
        assert bant['throat_diameter'][1] == float(
            motor_validator.limits['throat_diameter']['max'])
        assert bant['chamber_pressure'][1] == float(
            motor_validator.limits['chamber_pressure']['max'])

    @pytest.mark.parametrize('yol,govde,beklenen', [
        ('/api/analysis/wall-profile',
         {'chamber_pressure': 2.0e6, 'chamber_temperature': 3000.0,
          'chamber_diameter': 0.1, 'chamber_length': 0.5, 'burn_time': 10.0,
          'mdot_total': 1.0, 'material': 'steel', 'wall_thickness': 0.005,
          'throat_diameter': 0.03}, 422),
        ('/api/water-hammer',
         {'fluid': 'water', 'line_length_m': 5.0, 'line_id_mm': 0.02,
          'wall_thickness_mm': 1.5, 'working_pressure_bar': 25.0,
          'mdot_kg_s': 3.74}, 422),
        ('/analyze_thermal_safety',
         {'chamber_pressure': 20.0, 'chamber_temperature': 3000.0,
          'chamber_diameter': 100.0, 'chamber_length': 0.5,
          'burn_time': 10.0, 'mdot_total': 1.0, 'material': 'steel',
          'throat_diameter': 0.03}, 422),
    ])
    def test_kardes_uclarin_kurulu_kapilari_gerilemez(self, client, yol,
                                                      govde, beklenen):
        r = _post(client, yol, govde)
        assert r.status_code == beklenen, (
            '%s: HTTP %d' % (yol, r.status_code))


# ===========================================================================
# P7B1-20 — BAYRAK ALANI bool() İLE OKUNMAZ
# ===========================================================================
class TestP7B120BayrakOkumasi:
    """ÖLÇÜLDÜ (``once.json`` K6, ``/calculate``): ``include_plots='false'``
    -> 200 + 2.678.462 B + ``plots`` DOLU (13,03 s'lik tam figür seti);
    ``include_plots=False`` -> 200 + 1.481.231 B + ``plots`` YOK.
    ``bool('false')`` Python'da ``True``dur.
    """

    @pytest.mark.parametrize('ham,beklenen', [
        (False, False), (True, True), (0, False), (1, True),
        ('false', False), ('False', False), ('FALSE', False),
        ('no', False), ('off', False), ('0', False),
        ('true', True), ('yes', True), ('on', True), ('1', True),
        (' false ', False),
    ])
    def test_metin_bayraklar_adiyla_cozulur(self, ham, beklenen):
        from hrma.app import _request_flag
        deger, ret = _request_flag({'f': ham}, 'f', True)
        assert ret is None, '%r reddedildi' % (ham,)
        assert deger is beklenen, '%r -> %r' % (ham, deger)

    @pytest.mark.parametrize('ham', ['maybe', 'falsch', 'y', 2, -1, 1.5,
                                     [], {}])
    def test_cozulemeyen_bayrak_sessizce_bir_tarafa_dusmez(self, ham):
        from hrma.app import _request_flag
        # ``jsonify`` uygulama baglami ister; ret govdesi uc icinde uretilir.
        with app.test_request_context():
            _, ret = _request_flag({'f': ham}, 'f', True)
            assert ret is not None, '%r sessizce kabul edildi' % (ham,)
            govde, kod = ret
            assert kod == 422
            j = govde.get_json()
            assert j.get('error') == 'invalid_flag', j
            alan = (j.get('invalid_fields') or [{}])[0]
            assert alan.get('field') == 'f' and \
                alan.get('reason') == 'not_a_boolean', j

    @pytest.mark.parametrize('govde,beklenen', [
        ({}, True), ({'f': None}, True), ({'f': ''}, True),
    ])
    def test_verilmeyen_alan_varsayilana_doner(self, govde, beklenen):
        """``''`` ve ``null`` deponun kanonuna göre "verilmedi"dir."""
        from hrma.app import _request_flag
        deger, ret = _request_flag(govde, 'f', beklenen)
        assert ret is None and deger is beklenen

    def test_calculate_metin_false_figur_uretmez(self, client):
        """Ürün yolu: ölçülen 2.678.462 B'lik figür seti artık üretilmez."""
        r = _post(client, '/calculate', {
            'thrust': 1000, 'burn_time': 10, 'of_ratio': 6.0,
            'chamber_pressure': 20.0, 'fuel_type': 'htpb',
            'oxidizer_type': 'n2o', 'include_plots': 'false'})
        assert r.status_code == 200, r.get_data(as_text=True)[:200]
        j = r.get_json() or {}
        assert not j.get('plots'), 'figür seti üretildi'


# ===========================================================================
# P7B1-21 — ÖLÜ YARDIMCI
# ===========================================================================
class TestP7B121OluYardimci:
    """ÖLÇÜLDÜ (``once.json`` K7, AST): ``hrma/app.py:_mm_to_m`` çağrı
    sayısı 0 iken künyesi yürürlükten kalkmış "VARSAYILANA doner"
    sözleşmesini anlatıyordu; ``flight_vehicle.py`` aynı adı farklı imzayla
    (1 argüman) taşıyordu.
    """

    @staticmethod
    def _app_agaci():
        import hrma.app as A
        return ast.parse(io.open(A.__file__, encoding='utf-8').read())

    def test_olu_mm_to_m_govdesi_kalmadi(self):
        agac = self._app_agaci()
        tanimlar = [n.name for n in ast.walk(agac)
                    if isinstance(n, ast.FunctionDef)
                    and n.name.startswith('_mm_to_m')]
        assert '_mm_to_m' not in tanimlar, (
            'ölü sarmalayıcı geri geldi: %r' % (tanimlar,))
        # Ürünün GERÇEKTEN okuduğu iki sarmalayıcı yerinde durur.
        assert '_mm_to_m_read' in tanimlar and \
            '_mm_to_m_optional' in tanimlar, tanimlar

    def test_kalan_sarmalayicilar_merkeze_devreder(self):
        import inspect

        import hrma.app as A
        for ad in ('_mm_to_m_read', '_mm_to_m_optional'):
            kaynak = inspect.getsource(getattr(A, ad))
            assert 'input_guard.mm_to_m_read' in kaynak or \
                '_mm_to_m_read(' in kaynak, ad


# ===========================================================================
# S1-05 — İHRACAT AİLESİNDE HTTP SINIFI
# ===========================================================================
IHRACAT_AILESI = ['/api/export-stl', '/api/export-dxf', '/api/export-step',
                  '/api/export-drawings-pdf', '/api/export-complete-zip',
                  '/api/export-stl-zip', '/api/export-eng',
                  '/api/export-openrocket', '/api/export-simulation',
                  '/api/export-cad']


class TestS105AileSinifi:
    """ÖLÇÜLDÜ (``once.json`` K5; 10 uç x 4 gövde): BOŞ gövdede 9 uç 422
    (``missing_export_geometry``) verirken ``/api/export-stl`` HTTP **400**
    + yapısız gövde (``{'error': 'No motor data provided for export',
    'status': 'failed'}``) veriyordu — ne ailenin sınıfı ne ailenin şekli.
    """

    @pytest.mark.parametrize('yol', IHRACAT_AILESI)
    def test_bos_govde_ailede_tek_sinif(self, client, yol):
        r = _post(client, yol, {})
        assert r.status_code == 422, '%s: HTTP %d' % (yol, r.status_code)

    @pytest.mark.parametrize('yol', IHRACAT_AILESI)
    @pytest.mark.parametrize('tasiyici', [[1, 2, 3], 'abc'])
    def test_sozluk_olmayan_tasiyici_ailede_tek_sinif(self, client, yol,
                                                      tasiyici):
        r = _post(client, yol, {'motor_data': tasiyici})
        assert r.status_code == 422, '%s: HTTP %d' % (yol, r.status_code)

    def test_stl_ucu_ailenin_makine_okur_govdesini_verir(self, client):
        """400 + yapısız gövde yerine ailenin 422 sözleşmesi."""
        r = _post(client, '/api/export-stl', {})
        assert r.status_code == 422, r.status_code
        j = r.get_json() or {}
        assert (j.get('missing_fields') or j.get('invalid_fields')), j


# ===========================================================================
# ÖLÇÜM DOSYALARI — kanıt dizininin varlığı (rapor bağlanabilirliği)
# ===========================================================================
def test_olcum_kanitlari_yerinde():
    kok = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dizin = os.path.join(kok, 'docs', 'scofield', 'raw', 'OH', 'D29',
                         'app-api-uc', 'a2')
    for ad in ('once.json', 'sonra.json', 'olcum.py'):
        assert os.path.exists(kanit_gerekli(os.path.join(dizin, ad))), ad
    once = json.load(io.open(kanit_gerekli(os.path.join(dizin, 'once.json')),
                             encoding='utf-8'))
    # Taban ölçümünün kendisi de bir iddiadır: kusur ORADA vardı.
    assert once['K1_P16B4-20']['HAYALET']['kod'] == 200
    assert once['K3_P16B3-33']['kesirli_3.5']['kod'] == 200
    assert once['K5_S1-05']['/api/export-stl']['bos']['kod'] == 400
    assert math.isclose(once['K6_P7B1-20']['metin_false']['bayt'],
                        once['K6_P7B1-20']['yok']['bayt'])
